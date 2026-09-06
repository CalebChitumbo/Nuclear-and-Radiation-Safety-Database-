#!/usr/bin/env python3
"""
Importer: RPA "Daily Summary — All Inland Offices" workbook (.xlsx)
       -> seed/daily-screening-2026.seed.json

The Nuclear Safety, Security & Safeguards section keeps one sheet per inland
office, a row per calendar day, with the vehicles that post assessed that day.
A leading "Summary" sheet totals them by month, quarter and post.

This turns the per-post sheets into the section's daily log — one seeded
`dailyEntry` per post per day, on the same metric key the weekly report and
work plan output 1.3.12 already read — so the figures arrive as data the system
counts rather than as an opening balance nobody can drill into.

Decisions encoded here:
  * The per-post sheets are authoritative; the Summary sheet is only used to
    reconcile, month by month and post by post (a mismatch is reported, never
    silently absorbed).
  * A day with a blank cell was not reported and is skipped. A day the post
    wrote 0 on IS kept — "we screened nothing" is a reported figure.
  * Reporting weeks are assigned by the app (weekLabelForDate), not here, so a
    seeded entry and one an officer types on the same day land in the same week.

Usage:
  python3 scripts/convert-daily-summary-xlsx.py <Daily_Summary.xlsx> \
      [--out seed/daily-screening-2026.seed.json] \
      [--report docs/daily-screening-2026-import.md] \
      [--year 2026] [--imported-on YYYY-MM-DD]
"""

import argparse
import datetime
import json
import re
import sys
from collections import defaultdict
from datetime import date

try:
    import openpyxl
except ImportError:  # pragma: no cover - operator-facing message
    sys.exit(
        "openpyxl is missing, so nothing was read.\n\n"
        "The repo keeps a local Python environment for the workbook importers -\n"
        "macOS's own Python does not have openpyxl and should not be installed into.\n"
        "Use it:\n\n"
        "  .venv/bin/python "
        + " ".join(a if " " not in a else f'"{a}"' for a in sys.argv)
        + "\n\n"
        "or, if .venv is not there yet (one-off, ~20 seconds):\n\n"
        "  python3 -m venv .venv && .venv/bin/pip install openpyxl\n"
    )


SUMMARY_SHEET = "summary"

MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

# Rows of the Summary sheet that are roll-ups of the months above them.
SUMMARY_ROLLUPS = {"first quarter", "second quarter", "third quarter",
                   "fourth quarter", "total", "grand total"}


def source_name(path: str) -> str:
    """Workbook filename without the upload system's hash prefix."""
    return re.sub(r"^[0-9a-f]{6,}-", "", path.split("/")[-1])


def read_posts(wb, year: int):
    """Per-post daily counts: {post name: {date -> vehicles}}."""
    posts = {}
    problems = []
    for ws in wb.worksheets:
        if ws.title.strip().lower() == SUMMARY_SHEET:
            continue
        name = ws.title.strip()
        days = {}
        for row in ws.iter_rows(min_row=1, values_only=True):
            cells = list(row) + [None] * 3
            when, value = cells[1], cells[2]
            if not isinstance(when, datetime.datetime):
                continue  # month labels, the header, the per-month Total row
            if value is None:
                continue  # not reported that day
            if not isinstance(value, (int, float)):
                problems.append(f"{name} {when:%Y-%m-%d}: non-numeric {value!r}")
                continue
            iso = f"{when:%Y-%m-%d}"
            if when.year != year:
                problems.append(f"{name} {iso}: outside the plan year")
                continue
            if iso in days:
                problems.append(f"{name} {iso}: repeated row")
                continue
            days[iso] = int(value)
        if days:
            posts[name] = days
    return posts, problems


def read_summary(wb, posts):
    """The Summary sheet's month × post figures, for reconciliation."""
    ws = next(
        (w for w in wb.worksheets if w.title.strip().lower() == SUMMARY_SHEET),
        None,
    )
    if ws is None:
        return {}, {}

    header, rows = None, []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else c for c in row]
        labels = [str(c).strip() for c in cells]
        if header is None:
            if any(str(c).strip() in posts for c in cells):
                header = {
                    i: str(c).strip()
                    for i, c in enumerate(cells)
                    if str(c).strip() in posts
                }
            continue
        rows.append((labels, cells))

    by_month, by_post = {}, defaultdict(int)
    for labels, cells in rows:
        label = next((l for l in labels if l), "")
        key = label.lower()
        if label in MONTHS:
            by_month[label] = {
                post: int(cells[i] or 0) for i, post in header.items()
            }
        elif key == "total":
            for i, post in header.items():
                by_post[post] = int(cells[i] or 0)
    return by_month, dict(by_post)


def month_totals(days):
    """{month name: total} for one post's daily counts."""
    totals = defaultdict(int)
    for iso, value in days.items():
        totals[MONTHS[int(iso[5:7]) - 1]] += value
    return dict(totals)


def convert(path: str, year: int):
    wb = openpyxl.load_workbook(path, data_only=True)
    posts, problems = read_posts(wb, year)
    if not posts:
        sys.exit("No per-post sheets with daily counts were found")
    summary_months, summary_posts = read_summary(wb, posts)

    seed = {
        "source": source_name(path),
        "year": year,
        "posts": [
            {
                "border": name,
                "days": [[iso, days[iso]] for iso in sorted(days)],
            }
            for name, days in sorted(posts.items())
        ],
    }

    # ---- reconciliation, post by post and month by month
    checks = []
    for name, days in sorted(posts.items()):
        got = sum(days.values())
        expected = summary_posts.get(name)
        checks.append((name, expected, got, len(days)))
    month_checks = []
    for month in MONTHS:
        expected_row = summary_months.get(month)
        if expected_row is None:
            continue
        expected = sum(expected_row.values())
        got = sum(
            total
            for days in posts.values()
            for m, total in month_totals(days).items()
            if m == month
        )
        if expected or got:
            month_checks.append((month, expected, got))

    report = {
        "checks": checks,
        "month_checks": month_checks,
        "problems": problems,
        "entries": sum(len(d) for d in posts.values()),
        "total": sum(sum(d.values()) for d in posts.values()),
        "grand_total": sum(summary_posts.values()) if summary_posts else None,
    }
    return seed, report


def write_report(path, seed, report, imported_on):
    L = []
    year = seed["year"]
    L.append(f"# Daily screening import — {year} inland offices\n")
    L.append(
        f"Imported {imported_on} from `{seed['source']}` (the Nuclear Safety, "
        "Security & Safeguards section's daily summary workbook). Regenerate "
        "with:\n"
    )
    L.append(
        "```\npython3 scripts/convert-daily-summary-xlsx.py "
        "<Daily_Summary.xlsx>\n```\n"
    )
    L.append(
        "Each row becomes one daily `count` entry against **Vehicle Screening "
        "(units)** — the metric behind work plan output **1.3.12 Monitoring of "
        "illicit trafficking (ZRA Asycuda)** — tagged with the post it came "
        "from. The output's opening balance is therefore **zero**: the figures "
        "are counted from these entries, not carried in.\n"
    )

    L.append("## Result\n")
    L.append("| Metric | Count |")
    L.append("|---|---|")
    L.append(f"| Border posts | **{len(seed['posts'])}** |")
    L.append(f"| Daily entries | **{report['entries']:,}** |")
    L.append(f"| Vehicles assessed | **{report['total']:,}** |")
    L.append("")

    L.append("## Reconciled against the workbook's Summary sheet\n")
    L.append("| Post | Summary total | Imported | Days reported |")
    L.append("|---|---|---|---|")
    for name, expected, got, days in report["checks"]:
        if expected is None:
            mark = " — not on the Summary sheet"
            expected_text = "—"
        else:
            mark = "" if expected == got else f" ⚠ ({got - expected:+,})"
            expected_text = f"{expected:,}"
        L.append(f"| {name} | {expected_text} | {got:,}{mark} | {days} |")
    if report["grand_total"] is not None:
        mark = "" if report["grand_total"] == report["total"] else " ⚠"
        L.append(
            f"| **All posts** | **{report['grand_total']:,}** | "
            f"**{report['total']:,}**{mark} | {report['entries']} |"
        )
    L.append("")

    if report["month_checks"]:
        L.append("By month, across all posts:\n")
        L.append("| Month | Summary total | Imported |")
        L.append("|---|---|---|")
        for month, expected, got in report["month_checks"]:
            mark = "" if expected == got else f" ⚠ ({got - expected:+,})"
            L.append(f"| {month} | {expected:,} | {got:,}{mark} |")
        L.append("")

    L.append("## Quarters\n")
    L.append(
        "The workbook totals by CALENDAR quarter; the report totals by "
        "REPORTING WEEK, and a week counts to the quarter it starts in. Week "
        f"14 runs 30 Mar – 3 Apr {year}, so the three April days inside it are "
        "Q2 in the workbook and Q1 on the report. The year total is the same "
        "either way.\n"
    )

    if report["problems"]:
        L.append(f"## Rows not imported ({len(report['problems'])})\n")
        for p in report["problems"]:
            L.append(f"- {p}")
        L.append("")
    else:
        L.append(
            "Every dated row with a figure was imported; blank days are days "
            "the post did not report.\n"
        )

    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default="seed/daily-screening-2026.seed.json")
    ap.add_argument("--report", default="docs/daily-screening-2026-import.md")
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument(
        "--imported-on",
        default=date.today().isoformat(),
        help="Date recorded in the report header (default: today).",
    )
    args = ap.parse_args()

    seed, report = convert(args.xlsx, args.year)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(seed, fh, ensure_ascii=False, separators=(",", ":"))
        fh.write("\n")
    write_report(args.report, seed, report, args.imported_on)

    print(
        f"Wrote {args.out}: {len(seed['posts'])} posts, "
        f"{report['entries']} daily entries, {report['total']:,} vehicles"
    )
    for name, expected, got, _ in report["checks"]:
        if expected is None:
            print(f"  {name:<16} not on the Summary sheet   imported {got:>8,}")
        else:
            flag = "ok" if expected == got else f"MISMATCH ({got - expected:+,})"
            print(
                f"  {name:<16} summary {expected:>8,}  imported {got:>8,}  {flag}"
            )
    print(f"Report: {args.report}")
    if report["problems"]:
        print(f"  {len(report['problems'])} rows not imported — see the report")


if __name__ == "__main__":
    main()
