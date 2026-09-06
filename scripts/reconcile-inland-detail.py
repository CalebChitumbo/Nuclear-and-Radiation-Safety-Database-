#!/usr/bin/env python3
"""
Reconcile the inland offices' DETAIL workbooks against the daily summary.

The section keeps two things. The daily summary workbook (one figure per post
per day) is what output 1.3.12 counts, and it is already imported —
`seed/daily-screening-2026.seed.json`. Behind it sits a folder of monthly
workbooks per post, one sheet per day, one row per scanned truck: registration
or chassis number, the cargo typed into one of three columns, the transporter,
and the dose in nSv/h.

Nothing has ever checked that the second adds up to the first. This does: for
every day sheet it can read, it counts the truck rows and holds them against
the figure the summary reports for that post and day.

Four answers come out of it, and the difference between the middle two is the
whole point:

  * DAYS THAT AGREE — the figure is evidenced truck by truck.
  * DAYS WHERE THE SHEET IS BLANK — the workbook has that day's sheet and
    nobody filled it in, while the summary reports a figure for it. Not a
    disagreement: an unwritten page. Whole post-months look like this.
  * DAYS THAT GENUINELY DISAGREE — both sides carry rows and they do not
    match. Neither is assumed right; the difference is for the section to
    settle. These are the ones worth anybody's time.
  * DAYS WITH NO WORKBOOK AT ALL — the summary reports a figure and no book
    was sent, or it arrived as a PDF, or in a layout this cannot read.

Layouts
-------
Most post-months use one shared template, and only those are read:

    REG. NUMBER / CHASSIS | GOODS OF INTEREST | FOOD | OTHER |
    TRANSPORTER / DECLARANT | DOSE (nSv/h)

with a sheet per day named 1st, 2nd, 3rd … Anything else — Livingstone's
January–March plate lists, Chirundu's early "ASSESSEMENTS" books, Chingola's
March consignment layout, the Monthly Summary/Master Data books — is reported
as unread rather than guessed at. A wrong reading is worse than a gap.

Usage:
  python3 scripts/reconcile-inland-detail.py "INLAND DAILY ASSESSMENTS" \\
      [--seed seed/daily-screening-2026.seed.json] \\
      [--json docs/inland-detail-reconciliation.json] \\
      [--report docs/inland-detail-reconciliation.md]
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict

try:
    import openpyxl
except ImportError:  # pragma: no cover - operator-facing message
    sys.exit("openpyxl is required:  pip install openpyxl")


# The shared template's first row, normalised. A sheet whose header does not
# start this way is a different layout and is left alone.
TEMPLATE_HEAD = ("reg. number", "goods of interest")

MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11,
    "december": 12,
    # As the posts abbreviate them in file names.
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "jun": 6, "jul": 7, "aug": 8,
    "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12,
}

# Folder name -> the post name the summary workbook and the app use.
POST_ALIASES = {
    "kapiri- mposhi": "Kapiri Mposhi",
    "kapiri-mposhi": "Kapiri Mposhi",
}

DAY_SHEET = re.compile(r"^\s*(\d{1,2})\s*(st|nd|rd|th)\s*$", re.I)


def norm(v) -> str:
    return "" if v is None else str(v).strip()


def post_for(path: str, root: str) -> str:
    """The post a workbook belongs to: the first folder under the root."""
    rel = os.path.relpath(path, root)
    top = rel.split(os.sep)[0]
    return POST_ALIASES.get(top.lower(), top)


def month_for(path: str) -> int | None:
    """
    The month a workbook covers, from its own name or its folder's.

    Read from the file name first ("08 August.xlsx", "AUGUST NKD 2026.xlsx"),
    then the containing folder ("Chirundu/AUGUST 2026/August.xlsx"), so a book
    named only "August.xlsx" still resolves.
    """
    for part in (os.path.basename(path), os.path.basename(os.path.dirname(path))):
        words = re.findall(r"[a-z]+", part.lower())
        for w in words:
            if w in MONTHS:
                return MONTHS[w]
    return None


def is_template(ws) -> bool:
    row = next(ws.iter_rows(min_row=1, max_row=1, max_col=2, values_only=True), ())
    cells = [norm(c).lower() for c in row] + ["", ""]
    return cells[0].startswith(TEMPLATE_HEAD[0]) and cells[1].startswith(TEMPLATE_HEAD[1])


def count_trucks(ws) -> int:
    """
    Rows on a day sheet that record a scanned truck.

    A row counts when it carries a registration/chassis number. The tally block
    the officers typed to the right (commodity, count) lives in later columns
    and is not a truck, so only column A decides. Blank rows inside the block
    are skipped rather than treated as the end: the books have gaps.
    """
    trucks = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        reg = norm(row[0])
        if not reg:
            continue
        # The tally block sometimes bleeds a "0" or a total into column A.
        if reg in {"0", "TOTAL", "Total", "total"}:
            continue
        trucks += 1
    return trucks


def read_workbook(path: str, root: str):
    """Truck counts per date for one workbook, plus why sheets were skipped."""
    post = post_for(path, root)
    month = month_for(path)
    counts: dict[str, int] = {}
    skipped: list[str] = []

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # a corrupt or password-protected book
        return post, month, {}, [f"could not open: {type(exc).__name__}"]

    try:
        if month is None:
            return post, None, {}, ["no month in the file or folder name"]
        for name in wb.sheetnames:
            m = DAY_SHEET.match(name)
            if not m:
                continue
            day = int(m.group(1))
            ws = wb[name]
            if not is_template(ws):
                skipped.append(f"{name}: not the shared template")
                continue
            date = f"2026-{month:02d}-{day:02d}"
            counts[date] = count_trucks(ws)
    finally:
        wb.close()

    return post, month, counts, skipped


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("root", help="the INLAND DAILY ASSESSMENTS folder")
    ap.add_argument("--seed", default="seed/daily-screening-2026.seed.json")
    ap.add_argument("--json", dest="json_out")
    ap.add_argument("--report")
    args = ap.parse_args()

    summary: dict[str, dict[str, int]] = {}
    with open(args.seed, encoding="utf-8") as fh:
        for post in json.load(fh)["posts"]:
            summary[post["border"]] = {d: v for d, v in post["days"]}

    detail: dict[str, dict[str, int]] = defaultdict(dict)
    unread: list[tuple[str, str]] = []
    books = 0

    for dirpath, _dirs, files in os.walk(args.root):
        for name in sorted(files):
            if not name.endswith(".xlsx") or name.startswith("~$"):
                continue
            path = os.path.join(dirpath, name)
            if "Daily Summary" in path:
                continue
            books += 1
            post, _month, counts, skipped = read_workbook(path, args.root)
            rel = os.path.relpath(path, args.root)
            if not counts:
                unread.append((rel, skipped[0] if skipped else "no day sheets in the shared template"))
                continue
            for date, trucks in counts.items():
                # The same day in two books (a re-send) keeps the fuller one.
                detail[post][date] = max(detail[post].get(date, 0), trucks)

    lines: list[str] = []
    out = lines.append
    out("# Inland detail vs the daily summary\n")
    out(
        f"Read {books} monthly workbooks under `{args.root}`. The daily summary "
        f"(`{args.seed}`) is the figure output 1.3.12 counts; the detail books are "
        "the trucks behind it, one row each.\n"
    )

    agree = blank = disagree = no_book = only_detail = 0
    evidenced = summary_total = rows_total = 0
    per_post = []
    biggest: list[tuple[int, str, str, int, int]] = []

    for post in sorted(summary):
        s_days = summary[post]
        d_days = detail.get(post, {})
        s_total = sum(s_days.values())
        summary_total += s_total
        rows_total += sum(d_days.values())
        p_agree = p_blank = p_disagree = p_no_book = 0
        p_evidenced = 0
        for date, figure in s_days.items():
            if date not in d_days:
                p_no_book += 1
                continue
            trucks = d_days[date]
            if trucks == figure:
                p_agree += 1
                p_evidenced += figure
            elif trucks == 0:
                # The sheet is there and empty: an unwritten page, not a clash.
                p_blank += 1
            else:
                p_disagree += 1
                biggest.append((abs(trucks - figure), post, date, figure, trucks))
        for date in d_days:
            if date not in s_days:
                only_detail += 1
        agree += p_agree
        blank += p_blank
        disagree += p_disagree
        no_book += p_no_book
        evidenced += p_evidenced
        per_post.append(
            (post, len(s_days), s_total, p_agree, p_blank, p_disagree, p_no_book, p_evidenced)
        )

    out("## Per post\n")
    out("| Post | Days in summary | Vehicles | Evidenced | Sheet blank | Disagree | No workbook | Vehicles evidenced |")
    out("|---|---|---|---|---|---|---|---|")
    for post, days, total, ok, bl, bad, missing, ev in per_post:
        out(
            f"| {post} | {days} | {total:,} | {ok} | {bl} | {bad} | {missing} | "
            f"{ev:,} ({ev * 100 // total if total else 0}%) |"
        )
    out(
        f"| **All posts** | **{sum(p[1] for p in per_post)}** | "
        f"**{summary_total:,}** | **{agree}** | **{blank}** | **{disagree}** | "
        f"**{no_book}** | **{evidenced:,}** "
        f"({evidenced * 100 // summary_total if summary_total else 0}%) |\n"
    )
    out(
        f"{rows_total:,} truck rows were read in total across every day sheet, "
        f"against {summary_total:,} vehicles in the summary.\n"
    )

    if biggest:
        out("## The days that genuinely disagree\n")
        out(
            "Both sides carry rows here and they do not match — blank sheets are "
            "counted separately above and are not in this list. Neither figure is "
            "assumed right: the summary is what the report counts, the rows are "
            "what the post wrote down.\n"
        )
        out("| Post | Day | Summary says | Rows in the book | Difference |")
        out("|---|---|---|---|---|")
        for _d, post, date, figure, trucks in sorted(biggest, reverse=True)[:40]:
            out(f"| {post} | {date} | {figure:,} | {trucks:,} | {trucks - figure:+,} |")
        if len(biggest) > 40:
            out(f"\n… and {len(biggest) - 40} more, all in the JSON.\n")
        else:
            out("")

    if only_detail:
        out(
            f"{only_detail} day sheets hold trucks for days the summary does not "
            "report at all — the detail is ahead of the summary there.\n"
        )

    if unread:
        out("## Workbooks not read\n")
        out("Read only the shared template (REG. NUMBER / GOODS OF INTEREST / FOOD / OTHER / TRANSPORTER / DOSE, a sheet per day). A wrong reading would be worse than a gap.\n")
        out("| Workbook | Why |")
        out("|---|---|")
        for rel, why in sorted(unread):
            out(f"| `{rel}` | {why} |")
        out("")

    report = "\n".join(lines)
    if args.report:
        with open(args.report, "w", encoding="utf-8") as fh:
            fh.write(report)
        print(f"Report: {args.report}")
    else:
        print(report)

    if args.json_out:
        with open(args.json_out, "w", encoding="utf-8") as fh:
            json.dump(
                {
                    "detail": {p: dict(sorted(d.items())) for p, d in sorted(detail.items())},
                    "differences": [
                        {"post": p, "date": dt, "summary": f, "rows": t}
                        for _d, p, dt, f, t in sorted(biggest, reverse=True)
                    ],
                    "unread": [{"workbook": r, "why": w} for r, w in sorted(unread)],
                },
                fh,
                indent=2,
            )
            fh.write("\n")
        print(f"JSON: {args.json_out}")

    print(
        f"\n{books} workbooks · {agree} days evidenced · {blank} blank sheets · "
        f"{disagree} genuinely disagree · {no_book} days with no workbook · "
        f"{evidenced:,} of {summary_total:,} vehicles evidenced truck by truck"
    )


if __name__ == "__main__":
    main()
